from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth import get_user_model
from django.contrib import messages
from django.views.decorators.http import require_http_methods
from django.db.models import Q
from .forms import EmpleadoCreationForm, EmpleadoChangeForm

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image
from io import BytesIO
from django.http import HttpResponse
import os
from django.conf import settings

User = get_user_model()

def es_admin(user):
    return user.is_authenticated and user.rol == 'admin'

@login_required
@user_passes_test(es_admin)
@require_http_methods(["GET", "POST"])
def gestion_empleados(request):
    busqueda = request.GET.get('buscador_empleados', '').strip()

    # Crear empleado
    if request.method == 'POST' and 'crear' in request.POST:
        form_crear = EmpleadoCreationForm(request.POST)
        if form_crear.is_valid():
            form_crear.save()
            messages.success(request, 'Empleado creado correctamente.')
            return redirect('empleados:gestion_empleados')
        else:
            messages.error(request, 'Error al crear empleado. Revisa los datos.')
    else:
        form_crear = EmpleadoCreationForm()

    # Editar empleado
    if request.method == 'POST' and 'editar' in request.POST:
        empleado_id = request.POST.get('empleado_id')
        empleado = get_object_or_404(User, pk=empleado_id)
        form_editar = EmpleadoChangeForm(request.POST, instance=empleado)
        if form_editar.is_valid():
            form_editar.save()
            messages.success(request, 'Empleado actualizado correctamente.')
            return redirect('empleados:gestion_empleados')
        else:
            messages.error(request, 'Error al actualizar empleado.')
    else:
        form_editar = EmpleadoChangeForm()

    # Eliminar empleado
    if request.method == 'POST' and 'eliminar' in request.POST:
        empleado_id = request.POST.get('eliminar_id')
        empleado = get_object_or_404(User, pk=empleado_id)
        try:
            empleado.delete()
            messages.success(request, 'Empleado eliminado correctamente.')
        except Exception:
            messages.error(request, 'Error al eliminar empleado.')
        return redirect('empleados:gestion_empleados')

    # Usuarios filtrados (no clientes ni superusuarios)
    empleados = User.objects.exclude(rol=User.CLIENTE).exclude(is_superuser=True)

    # Filtro con búsqueda, o mostrar todos si vacío/no resultados
    if busqueda:
        empleados_filtrados = empleados.filter(
            Q(username__icontains=busqueda) |
            Q(email__icontains=busqueda) |
            Q(rol__icontains=busqueda)
        )
        if empleados_filtrados.exists():
            empleados = empleados_filtrados
        else:
            # No hay coincidencias, mostrar todos
            pass
    else:
        # Búsqueda vacía, mostrar todos (ya asignado en empleados)
        pass

    return render(request, 'empleados/gestion_empleados.html', {
        'empleados': empleados,
        'form_crear': form_crear,
        'form_editar': form_editar,
        'busqueda': busqueda,
    })


def exportar_empleados_excel(request):
    wb = openpyxl.Workbook()

    # Ruta ingreso logo
    logo_path = os.path.join(settings.BASE_DIR, 'app/static/img/logos/logo_asvg.png')
    if not os.path.exists(logo_path):
        raise FileNotFoundError(f"Logo no encontrado en: {logo_path}")

    ws = wb.active
    ws.title = "Empleados"

    # Cabecera con nombre empresa y logo
    ws.merge_cells('A1:B1')
    ws['A1'] = "CONTACTOS JLARA"
    ws['A1'].font = Font(size=22, bold=True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    img = Image(logo_path)
    img.width = 100
    img.height = 70
    ws.column_dimensions['C'].width = 18
    ws.row_dimensions[1].height = 50
    img.anchor = 'C1'
    ws.add_image(img)

    empleados = User.objects.exclude(rol=User.CLIENTE).exclude(is_superuser=True)
    cantidad_empleados = empleados.count()

    ws['A3'] = f"Cantidad total de empleados: {cantidad_empleados}"
    ws['A3'].font = Font(size=14, bold=True)

    headers = ["Username", "Email", "Rol"]
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, size=13)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        side = Side(border_style='thin', color='000000')
        cell.border = Border(left=side, right=side, top=side, bottom=side)

    for row_num, empleado in enumerate(empleados, start=6):
        valores = [
            empleado.username,
            empleado.email,
            empleado.rol
        ]
        for col_num, valor in enumerate(valores, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = valor
            cell.alignment = Alignment(horizontal='center', vertical='center')
            side = Side(border_style='thin', color='000000')
            cell.border = Border(left=side, right=side, top=side, bottom=side)

    # Ajustar ancho columnas
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 15
    ws.row_dimensions[5].height = 30
    for row in range(6, 6 + cantidad_empleados):
        ws.row_dimensions[row].height = 20

    # Preparar y enviar archivo
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=empleados_contacto.xlsx'
    with BytesIO() as b:
        wb.save(b)
        response.write(b.getvalue())

    return response