from django.shortcuts import render, redirect
from django.contrib import messages
from .models import Productos
from .forms import ProductoForm
from categorias.models import Categorias
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from io import BytesIO
from django.http import HttpResponse
from .models import Productos
import os
from django.conf import settings
from openpyxl.utils import get_column_letter

def productos(request):
    filtro = request.GET.get('filtro', 'activos')

    if filtro == 'eliminados':
        productos = Productos.objects.filter(status=False)
    else:
        productos = Productos.objects.filter(status=True)

    categorias = Categorias.objects.filter(status=True)

    formulario = ProductoForm()

    if request.method == 'POST':
        if 'eliminar_id' in request.POST:
            try:
                producto = Productos.objects.get(id_producto_PK=request.POST.get('eliminar_id'))
                producto.status = False
                producto.save()
                messages.success(request, "Producto eliminado correctamente.")
            except Exception as e:
                messages.error(request, "Error: " + str(e))
            return redirect(f"{request.path}?filtro={filtro}")

        if 'activar_id' in request.POST:
            try:
                producto = Productos.objects.get(id_producto_PK=request.POST.get('activar_id'))
                producto.status = True
                producto.save()
                messages.success(request, "Producto reactivado correctamente.")
            except Exception as e:
                messages.error(request, "Error al reactivar: " + str(e))
            return redirect(f"{request.path}?filtro={filtro}")

        producto_id = request.POST.get('producto_id')

        if producto_id:
            producto = Productos.objects.get(id_producto_PK=producto_id)
            formulario = ProductoForm(request.POST, request.FILES, instance=producto)
            if formulario.is_valid():
                producto_actualizado = formulario.save(commit=False)
                producto_actualizado.status = True  
                producto_actualizado.save()
                messages.success(request, "Producto actualizado con éxito.")
            else:
                messages.error(request, "Corrige los errores en el formulario.")
            return redirect(f"{request.path}?filtro={filtro}")


        if not producto_id:
            formulario = ProductoForm(request.POST, request.FILES)
            if formulario.is_valid():
                producto_nuevo = formulario.save(commit=False)
                producto_nuevo.status = True 
                producto_nuevo.save()
                messages.success(request, "Producto creado con éxito.")
            else:
                messages.error(request, "Corrige los errores en el formulario.")
            return redirect(f"{request.path}?filtro={filtro}")

    return render(request, 'productos/productos.html', {
        'productos': productos,
        'formulario': formulario,
        'filtro': filtro,
        'categorias': categorias,
    })

def exportar_excel(request):
    wb = openpyxl.Workbook()
    
    # Ruta del logo
    logo_path = os.path.join(settings.BASE_DIR, 'app/static/img/logos/logo_asvg.png')
    if not os.path.exists(logo_path):
        raise FileNotFoundError(f"Logo no encontrado en: {logo_path}")

    # Hoja 1 para productos activos
    ws_activos = wb.active
    ws_activos.title = "Productos Activos"

    # Configuración del encabezado
    ws_activos.merge_cells('B1:E1')
    ws_activos['B1'] = "CONTACTOS JLARA"
    ws_activos['B1'].font = Font(size=22, bold=True)
    ws_activos['B1'].alignment = Alignment(horizontal='center', vertical='center')

    # Insertar logo en el encabezado
    img = Image(logo_path)
    img.width = 140
    img.height = 70
    ws_activos.column_dimensions['F'].width = 18  
    ws_activos.row_dimensions[1].height = 50
    img.anchor = 'F1'
    ws_activos.add_image(img)

    # Consulta productos activos
    activo_qs = Productos.objects.filter(status=True)
    cantidad_activos = activo_qs.count()

    # Insertar imagen izquierda
    imagen_izquierda_path = os.path.join(settings.BASE_DIR, 'app/static/img/logos/logo.png')
    imagen_izquierda = Image(imagen_izquierda_path)
    imagen_izquierda.width = 70
    imagen_izquierda.height = 80
    imagen_izquierda.anchor = 'A1'
    ws_activos.add_image(imagen_izquierda)

    # Cantidad de productos
    ws_activos['A3'] = f"Cantidad de productos existentes: {cantidad_activos}"
    ws_activos['A3'].font = Font(size=14, bold=True)

    # Encabezados de tabla
    headers = ["ID", "Nombre", "Descripción", "Precio", "Cantidad", "Categoría"]
    for col_num, header in enumerate(headers, 1):
        cell = ws_activos.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, size=13)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        side = Side(border_style='thin', color='000000')
        cell.border = Border(left=side, right=side, top=side, bottom=side)

    # Rellenar datos de productos activos
    for row_num, p in enumerate(activo_qs, start=6):
        valores = [
            p.id_producto_PK, 
            p.nombre, 
            p.descripcion, 
            float(p.precio),
            int(p.cantidad),
            p.categoria.descripcion if p.categoria else "Sin categoría"
        ]
        for col_num, valor in enumerate(valores, 1):
            cell = ws_activos.cell(row=row_num, column=col_num)
            cell.value = valor
            cell.alignment = Alignment(horizontal='center', vertical='center')
            side = Side(border_style='thin', color='000000')
            cell.border = Border(left=side, right=side, top=side, bottom=side)
    # Ajustes de ancho de columnas y alto de filas
    for col in range(1, len(headers) + 1):
        columna = get_column_letter(col)
        ws_activos.column_dimensions[columna].width = 25

    ws_activos.column_dimensions['A'].width = 10  
    ws_activos.column_dimensions['C'].width = 40  
    ws_activos.column_dimensions['D'].width = 15  
    ws_activos.column_dimensions['E'].width = 15  
    ws_activos.column_dimensions['D'].width = 20  

    ws_activos.row_dimensions[5].height = 30  
    for row in range(6, 6 + cantidad_activos):
        ws_activos.row_dimensions[row].height = 20  

    # Hoja 2 para productos inactivos
    ws_inactivos = wb.create_sheet(title="Productos Inactivos")

    # Encabezado hoja inactivos
    ws_inactivos.merge_cells('B1:E1')
    ws_inactivos['B1'] = "CONTACTOS JLARA"
    ws_inactivos['B1'].font = Font(size=22, bold=True)
    ws_inactivos['B1'].alignment = Alignment(horizontal='center', vertical='center')

    # Insertar logo en el encabezado de inactivos
    img_inactivos = Image(logo_path)
    img_inactivos.width = 130
    img_inactivos.height = 70
    ws_inactivos.column_dimensions['F'].width = 18  
    ws_inactivos.row_dimensions[1].height = 50
    img_inactivos.anchor = 'F1'
    ws_inactivos.add_image(img_inactivos)

    # Imagen izquierda en hoja inactivos
    imagen_izquierda_inactivos_path = os.path.join(settings.BASE_DIR, 'app/static/img/logos/logo.png')
    imagen_izquierda_inactivos = Image(imagen_izquierda_inactivos_path)
    imagen_izquierda_inactivos.width = 70
    imagen_izquierda_inactivos.height = 80
    imagen_izquierda_inactivos.anchor = 'A1'
    ws_inactivos.add_image(imagen_izquierda_inactivos)

    # Consulta productos inactivos
    inactivos_qs = Productos.objects.filter(status=False)
    cantidad_inactivos = inactivos_qs.count()

    # Cantidad de productos inactivos
    ws_inactivos['A3'] = f"Cantidad de productos existentes: {cantidad_inactivos}"
    ws_inactivos['A3'].font = Font(size=14, bold=True)

    # Encabezados de tabla para inactivos
    for col_num, header in enumerate(headers, 1):
        cell = ws_inactivos.cell(row=5, column=col_num)
        cell.value = header
        cell.font = Font(bold=True, size=13)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        side = Side(border_style='thin', color='000000')
        cell.border = Border(left=side, right=side, top=side, bottom=side)

    # Rellenar datos de productos inactivos
    for row_num, p in enumerate(inactivos_qs, start=6):
        valores = [
            p.id_producto_PK, 
            p.nombre, 
            p.descripcion, 
            float(p.precio),
            int(p.cantidad),
            p.categoria.descripcion if p.categoria else "Sin categoría"
        ]
        for col_num, valor in enumerate(valores, 1):
            cell = ws_inactivos.cell(row=row_num, column=col_num)
            cell.value = valor
            cell.alignment = Alignment(horizontal='center', vertical='center')
            side = Side(border_style='thin', color='000000')
            cell.border = Border(left=side, right=side, top=side, bottom=side)

    # Ajustes de ancho de columnas y alto de filas en hoja inactivos
    for col in range(1, len(headers) + 1):
        columna = get_column_letter(col)
        ws_inactivos.column_dimensions[columna].width = 25

    ws_inactivos.column_dimensions['A'].width = 10  
    ws_inactivos.column_dimensions['C'].width = 40  
    ws_inactivos.column_dimensions['D'].width = 15  
    ws_inactivos.column_dimensions['E'].width = 15  
    ws_inactivos.column_dimensions['F'].width = 20  

    ws_inactivos.row_dimensions[5].height = 30  
    for row in range(6, 6 + cantidad_inactivos):
        ws_inactivos.row_dimensions[row].height = 20  

    # Guardar libro en memoria y enviar como respuesta descargable
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename=productos_contacto.xlsx'
    with BytesIO() as b:
        wb.save(b)
        response.write(b.getvalue())

    return response
